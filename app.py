import io
import random
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from difflib import get_close_matches

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

st.set_page_config(page_title="Jadwal Angkatan 15", page_icon="✦", layout="wide")
ROSTER = {"Aliyah":"F", "Soma":"M", "Syamsul":"M", "Ferrel":"M", "Kezia":"F", "Alam":"M", "Bagus":"M", "Retno":"F", "Rachel":"F", "Irpan":"M", "Farez":"M"}
NAMES = list(ROSTER)
DAY_NAMES = ("Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu")
MONTHS = ("Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember")
WEEKDAY_WORDS = {"senin", "selasa", "rabu", "kamis", "jumat", "sabtu", "minggu"}
DOCTOR_WORDS = {"drg", "dr", "sp", "mmf", "mf", "subsp", "comf", "tr", "tm", "mars", "ph", "d", *WEEKDAY_WORDS}


def default_cohort_config():
    """Default only. Admin can add, remove, rename, or change each row before parsing."""
    return pd.DataFrame([
        {"Kolom": "a12", "Label angkatan": "Angkatan 12", "Residen per hari": 1, "Aktif": True},
        {"Kolom": "a13", "Label angkatan": "Angkatan 13", "Residen per hari": 2, "Aktif": True},
        {"Kolom": "a14", "Label angkatan": "Angkatan 14", "Residen per hari": 3, "Aktif": True},
        {"Kolom": "a15", "Label angkatan": "Angkatan 15", "Residen per hari": 4, "Aktif": True},
        {"Kolom": "a16", "Label angkatan": "Angkatan 16", "Residen per hari": 5, "Aktif": True},
        {"Kolom": "a17", "Label angkatan": "Angkatan 17", "Residen per hari": 5, "Aktif": True},
    ])


def roster_words(value):
    return re.findall(r"[A-Za-zÀ-ÿ]+(?:[-'][A-Za-zÀ-ÿ]+)?", str(value))


def clean_sentence(value):
    return re.sub(r"\s+", " ", str(value)).strip(" ,;")


def split_roster_blocks(text):
    """Split a Word/Google Docs table copy by its repeated weekday header."""
    header = re.compile(r"(?m)^(?=(?:(?:Senin|Selasa|Rabu|Kamis|Jumat|Sabtu|Minggu)(?:\s+|$))+$)")
    points = [match.start() for match in header.finditer(text)]
    if not points:
        return [text]
    return [text[start:end] for start, end in zip(points, points[1:] + [len(text)])]


def parse_pasted_roster(text, config):
    """Parse a visual table pasted as text without requiring a CSV upload.

    Word often wraps cells at different points. The parser first reads the weekly
    date block and then reconstructs each cohort from its expected member count.
    It also learns cohort membership from clean blocks, allowing it to recover
    rows whose columns were interleaved during copying.
    """
    active = config.copy()
    active["Kolom"] = active["Kolom"].astype(str).str.strip().str.lower()
    active["Residen per hari"] = pd.to_numeric(active["Residen per hari"], errors="coerce").fillna(0).astype(int)
    active = active[(active["Aktif"] == True) & active["Kolom"].ne("") & (active["Residen per hari"] > 0)]
    definitions = [(row["Kolom"], int(row["Residen per hari"])) for _, row in active.iterrows()]
    if not definitions:
        return pd.DataFrame(), ["Tambahkan minimal satu angkatan aktif dengan jumlah residen per hari lebih dari nol."], []

    blocks, warnings, skipped = [], [], []
    for block_number, block in enumerate(split_roster_blocks(text), start=1):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        dates = re.findall(r"\b\d{2}/\d{2}/\d{4}\b", block)
        if not dates:
            continue
        column_count = len(dates)
        date_line = max(index for index, line in enumerate(lines) if re.search(r"\d{2}/\d{2}/\d{4}", line))
        operational_start = None
        for index, line in enumerate(lines[date_line + 1:], start=date_line + 1):
            words = roster_words(line)
            if len(words) == column_count and all(word.lower() not in DOCTOR_WORDS for word in words):
                operational_start = index
                break
        if operational_start is None:
            skipped.extend(dates)
            warnings.append(f"Blok {block_number}: ditemukan {column_count} tanggal, tetapi baris Pilot tidak memiliki {column_count} nama. Blok tidak dipetakan agar tidak salah.")
            continue

        stream = roster_words(" ".join(lines[operational_start:]))
        required = column_count * (2 + sum(size for _, size in definitions))
        if len(stream) < required:
            skipped.extend(dates)
            warnings.append(f"Blok {block_number}: roster hanya memiliki {len(stream)} nama, sedangkan konfigurasi membutuhkan {required}. Blok tidak dipetakan.")
            continue

        doctors_text = " ".join(lines[date_line + 1:operational_start])
        doctor_parts = [clean_sentence(item) for item in re.split(r"(?i)(?=(?:dr\.\s*)?drg\.)", doctors_text) if re.search(r"(?i)(?:dr\.\s*)?drg\.", item)]
        provisional, position = {}, 2 * column_count
        for code, size in definitions:
            provisional[code] = stream[position:position + column_count * size]
            position += column_count * size
        blocks.append({
            "dates": dates,
            "count": column_count,
            "pilot": stream[:column_count],
            "copilot": stream[column_count:2 * column_count],
            "tail": stream[2 * column_count:],
            "provisional": provisional,
            "doctors": doctor_parts,
        })

    # Establish cohort identity from the first positional pass. This corrects the
    # common Word-copy issue where the A14/A15 visual cells appear interleaved.
    votes = defaultdict(Counter)
    display_names = {}
    for block in blocks:
        for code, people in block["provisional"].items():
            for person in people:
                key = person.lower()
                votes[key][code] += 1
                display_names.setdefault(key, person)
    owner = {person: counts.most_common(1)[0][0] for person, counts in votes.items()}
    known_names = list(display_names)

    rows = []
    for block in blocks:
        count = block["count"]
        corrected_tail = []
        for person in block["tail"]:
            key = person.lower()
            if key not in owner:
                close = get_close_matches(key, known_names, n=1, cutoff=.83)
                if close:
                    warnings.append(f"Nama `{person}` dibaca sebagai `{display_names[close[0]]}` karena ejaan sangat mirip.")
                    person, key = display_names[close[0]], close[0]
            corrected_tail.append(person)
        groups = {}
        for code, size in definitions:
            recovered = [person for person in corrected_tail if owner.get(person.lower()) == code]
            expected = count * size
            groups[code] = recovered if len(recovered) == expected else block["provisional"][code]
            if len(recovered) not in (0, expected):
                warnings.append(f"{block['dates'][0]}: {code} terbaca {len(recovered)}/{expected} nama setelah pemulihan; gunakan hasil preview untuk koreksi.")
        for index, raw_date in enumerate(block["dates"]):
            item = {
                "Tanggal": datetime.strptime(raw_date, "%d/%m/%Y").date().isoformat(),
                "DPJP": block["doctors"][index] if len(block["doctors"]) == count else "",
                "Pilot": block["pilot"][index],
                "Co-pilot": block["copilot"][index],
            }
            for code, size in definitions:
                people = groups[code][index * size:(index + 1) * size]
                item[code] = ", ".join(people)
            rows.append(item)
        if len(block["doctors"]) != count:
            warnings.append(f"{block['dates'][0]}: DPJP tidak dapat dipisahkan otomatis. Isi kolom DPJP pada tabel preview bila diperlukan.")

    if skipped:
        pretty = ", ".join(datetime.strptime(value, "%d/%m/%Y").strftime("%d %b") for value in skipped)
        warnings.append(f"Tanggal yang belum masuk preview: {pretty}.")
    return pd.DataFrame(rows), list(dict.fromkeys(warnings)), skipped


def render_roster_intake():
    st.markdown("<div class='masthead'><div class='service-line'>DEPARTEMEN BEDAH MULUT & MAKSILOFASIAL</div><h1>Pembagian Jaga</h1><p>Tempel tabel roster apa adanya. Sistem memetakan tanggal dan angkatan terlebih dahulu, lalu admin dapat mengoreksi hasil sebelum pembagian dibuat.</p></div>", unsafe_allow_html=True)
    st.markdown("<div class='panel'><b>1. Konfigurasi angkatan</b><br><span style='color:#60717d'>Kolom dan jumlah residen per hari tidak dikunci. Tambah, hapus, atau ubah label sebelum membaca roster.</span></div>", unsafe_allow_html=True)
    if "cohort_config" not in st.session_state:
        st.session_state.cohort_config = default_cohort_config()
    config = st.data_editor(
        st.session_state.cohort_config,
        num_rows="dynamic",
        hide_index=True,
        use_container_width=True,
        column_config={"Residen per hari": st.column_config.NumberColumn(min_value=1, step=1), "Aktif": st.column_config.CheckboxColumn()},
        key="cohort_config_editor",
    )
    st.session_state.cohort_config = config

    st.markdown("<div class='panel'><b>2. Tempel roster</b><br><span style='color:#60717d'>Tidak perlu unggah CSV. Paste langsung dari Word, Google Docs, atau tabel sumber.</span></div>", unsafe_allow_html=True)
    pasted = st.text_area("Roster yang ditempel", value=st.session_state.get("pasted_roster", ""), height=250, placeholder="Tempel seluruh tabel roster di sini…", key="pasted_roster_input")
    if st.button("Petakan roster", type="primary", use_container_width=False):
        st.session_state.pasted_roster = pasted
        parsed, warnings, skipped = parse_pasted_roster(pasted, config)
        st.session_state.parsed_roster = parsed
        st.session_state.roster_warnings = warnings
        st.session_state.roster_skipped = skipped

    parsed = st.session_state.get("parsed_roster")
    if parsed is None:
        st.info("Mulai dengan menempel roster, lalu cek hasil pemetaan. Pembagian klinis belum dibuat sebelum roster dinyatakan benar.")
        return
    warnings = st.session_state.get("roster_warnings", [])
    for warning in warnings:
        st.warning(warning)
    if parsed.empty:
        st.error("Belum ada tanggal yang dapat dipetakan. Cek kembali struktur paste dan jumlah residen per angkatan.")
        return

    labels = dict(zip(config["Kolom"].astype(str).str.lower(), config["Label angkatan"].astype(str)))
    shown = parsed.rename(columns=labels)
    st.markdown("<div class='panel'><b>3. Verifikasi dan koreksi</b><br><span style='color:#60717d'>Inilah sumber untuk pembagian berikutnya. Koreksi langsung di tabel bila nama atau DPJP terpotong saat copy-paste.</span></div>", unsafe_allow_html=True)
    edited = st.data_editor(shown, num_rows="dynamic", hide_index=True, use_container_width=True, height=520, key="parsed_roster_editor")
    reverse_labels = {label: code for code, label in labels.items()}
    st.session_state.parsed_roster = edited.rename(columns=reverse_labels)
    metrics = st.columns(3)
    metrics[0].metric("Tanggal terbaca", len(edited))
    metrics[1].metric("Angkatan aktif", len(labels))
    metrics[2].metric("Butuh koreksi", len(st.session_state.get("roster_skipped", [])))
    st.caption("Setelah tabel ini rapi, tahap berikutnya adalah mengaktifkan pembagian Post-op, Pre-op, dan IGD dari roster yang telah diverifikasi. Tidak ada data yang dikirim ke database atau membutuhkan unggahan CSV.")


def init_state():
    st.session_state.setdefault("unavailable", {name: set() for name in NAMES})
    st.session_state.setdefault("forbidden", set())
    st.session_state.setdefault("schedule", None)
    st.session_state.setdefault("summary", None)
    st.session_state.setdefault("notices", [])


def longest_streak(day_sets, name):
    best = current = 0
    for team in day_sets:
        current = current + 1 if name in team else 0
        best = max(best, current)
    return best


def targets_for(slots, rng):
    base, extra = divmod(slots, len(NAMES))
    order = NAMES[:]
    rng.shuffle(order)
    return {name: base + int(index < extra) for index, name in enumerate(order)}


def sunday_targets(days, quota, totals, rng):
    """Hard Sunday fairness: each person gets either floor or ceil share."""
    slots = sum(day.weekday() == 6 for day in days) * quota
    base, extra = divmod(slots, len(NAMES))
    order = NAMES[:]
    rng.shuffle(order)
    order.sort(key=lambda name: totals[name], reverse=True)
    result = {name: base for name in NAMES}
    for name in order[:extra]:
        result[name] += 1
    return result if all(result[name] <= totals[name] for name in NAMES) else None


def forbidden_with(picked, candidate, forbidden):
    return any(tuple(sorted((candidate, other))) in forbidden for other in picked)


def make_fair_jaga(days, quota, unavailable, forbidden, rng):
    if quota == 0:
        return [set() for _ in days], [], {name: 0 for name in NAMES}
    total_slots = len(days) * quota
    max_without_three = (len(days) // 3) * 2 + min(len(days) % 3, 2)
    enforce_streak = total_slots <= len(NAMES) * max_without_three
    notices = [] if enforce_streak else ["Batas dua Jaga berturut-turut dilonggarkan karena kuota terlalu tinggi; fairness total dan Minggu tetap dikunci."]
    for _ in range(1500):
        totals = targets_for(total_slots, rng)
        sunday = sunday_targets(days, quota, totals, rng)
        if sunday is None:
            continue
        remaining, sunday_left = Counter(totals), Counter(sunday)
        teams, failed = [], False
        for current in days:
            key, is_sunday = current.isoformat(), current.weekday() == 6
            yesterday = teams[-1] if teams else set()
            two_days_ago = teams[-2] if len(teams) > 1 else set()

            def eligible(name):
                if remaining[name] <= 0 or key in unavailable[name]:
                    return False
                if enforce_streak and name in yesterday and name in two_days_ago:
                    return False
                # Preserve Sunday share, so a Sunday target cannot be consumed on a weekday.
                return sunday_left[name] > 0 if is_sunday else remaining[name] > sunday_left[name]

            picked = []
            for _slot in range(quota):
                candidates = [name for name in NAMES if name not in picked and eligible(name) and not forbidden_with(picked, name, forbidden)]
                if not candidates:
                    failed = True
                    break

                def score(name):
                    gentle_pair_bias = 0.0
                    if {"Ferrel", "Alam"} & set(picked) and name in {"Ferrel", "Alam"} and rng.random() < .42:
                        gentle_pair_bias = .35
                    return (remaining[name], sunday_left[name] if is_sunday else 0, int(name not in yesterday), int(name not in two_days_ago), gentle_pair_bias, rng.random())
                picked.append(max(candidates, key=score))
            if failed:
                break
            for name in picked:
                remaining[name] -= 1
                if is_sunday:
                    sunday_left[name] -= 1
            teams.append(set(picked))
        if not failed and not any(remaining.values()) and not any(sunday_left.values()):
            return teams, notices, sunday
    return None, ["Tidak ada jadwal yang memenuhi fairness total, fairness Minggu, dan semua request. Ubah request atau kuota."], None


def choose_role(candidates, quota, counts, rng):
    if quota == 0:
        return []
    if len(candidates) < quota:
        return []
    return sorted(candidates, key=lambda name: (counts[name], rng.random()))[:quota]


def make_schedule(days, quotas, doru, unavailable, forbidden, seed):
    rng = random.Random(seed)
    teams, notices, _ = make_fair_jaga(days, quotas["Jaga"], unavailable, forbidden, rng)
    if teams is None:
        return None, None, notices
    counts = {role: Counter() for role in quotas}
    sunday, rows = Counter(), []
    for index, current in enumerate(days):
        jaga = sorted(teams[index]); used = set(jaga)
        review = choose_role([name for name in NAMES if name not in used and name not in doru], quotas["Review"], counts["Review"], rng)
        if len(review) != quotas["Review"]:
            return None, None, [f"{current:%d %b}: kuota Review tidak dapat dipenuhi."]
        used.update(review)
        erm = choose_role([name for name in NAMES if name not in used and name not in doru], quotas["ERM"], counts["ERM"], rng)
        if len(erm) != quotas["ERM"]:
            return None, None, [f"{current:%d %b}: kuota ERM tidak dapat dipenuhi."]
        for role, people in (("Jaga", jaga), ("Review", review), ("ERM", erm)):
            counts[role].update(people)
        if current.weekday() == 6:
            sunday.update(jaga)
        rows.append({"date":current.isoformat(), "Tanggal":current.strftime("%a, %d %b %Y"), "jaga":", ".join(jaga), "review":", ".join(review), "erm":", ".join(erm), "Jaga":", ".join(jaga), "Review":", ".join(review), "ERM":", ".join(erm)})
    total = [counts["Jaga"][name] for name in NAMES]
    sunday_counts = [sunday[name] for name in NAMES]
    streaks = {name: longest_streak(teams, name) for name in NAMES}
    if max(total) - min(total) > 1 or max(sunday_counts) - min(sunday_counts) > 1:
        return None, None, ["Validasi fairness gagal; tidak ada jadwal diekspor."]
    if max(streaks.values()) > 2 and not notices:
        return None, None, ["Validasi Jaga berturut-turut gagal; tidak ada jadwal diekspor."]
    summary = pd.DataFrame({"Nama":NAMES, "Jaga":total, "Jaga Minggu":sunday_counts, "Jaga Berturut Terpanjang":[streaks[name] for name in NAMES], "Review":[counts["Review"][name] for name in NAMES], "ERM":[counts["ERM"][name] for name in NAMES]})
    return pd.DataFrame(rows), summary, notices


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr(); node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd"); props.append(node)
    node.set(qn("w:fill"), color)


def bordered(cell):
    props = cell._tc.get_or_add_tcPr(); borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); props.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "6"); node.set(qn("w:color"), "1F2937")


def write_cell(cell, text, size=8, bold=False):
    cell.text = ""; p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text); run.font.name = "Times New Roman"; run.font.size = Pt(size); run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; bordered(cell)


def weeks_between(start, end):
    current = start - timedelta(days=start.weekday()); finish = end + timedelta(days=6-end.weekday()); result = []
    while current <= finish:
        result.append([current + timedelta(days=offset) if start <= current + timedelta(days=offset) <= end else None for offset in range(7)])
        current += timedelta(days=7)
    return result


def schedule_docx(schedule, summary, start, end):
    doc = Document(); section = doc.sections[0]; section.orientation = WD_ORIENT.LANDSCAPE; section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.5); section.right_margin = Cm(.6); section.top_margin = Cm(1.4); section.bottom_margin = Cm(.6)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("JADWAL JAGA REVIEW DAN ERM ANGKATAN 15"); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; label = f"{MONTHS[start.month-1].upper()} {start.year}" if (start.month,start.year)==(end.month,end.year) else f"{start:%d/%m/%Y} sampai {end:%d/%m/%Y}"; r = p.add_run(label); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(12)
    assignments = {date.fromisoformat(row.date):row for row in schedule.itertuples()}
    for week_index, week in enumerate(weeks_between(start,end)):
        if week_index and week_index % 2 == 0: doc.add_page_break()
        table = doc.add_table(rows=5, cols=7); table.autofit = False
        for column, current in enumerate(week):
            color = "E7B008" if column != 6 else "E05252"; shade(table.cell(0,column),color); write_cell(table.cell(0,column),DAY_NAMES[column],10,True); shade(table.cell(1,column),color); write_cell(table.cell(1,column),current.strftime("%d/%m/%Y") if current else "",9,True)
            for row_index, role in enumerate(("jaga","review","erm"),2):
                cell = table.cell(row_index,column)
                if current:
                    if row_index == 2: shade(cell,"F3F4F6")
                    write_cell(cell,f"{role.title()}\n{getattr(assignments[current],role)}",7.4)
                else: write_cell(cell,"",8)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    heading = doc.add_heading("Rekap Fairness Jaga",level=2)
    for run in heading.runs: run.font.name="Times New Roman"; run.font.color.rgb=RGBColor(0,0,0)
    note=doc.add_paragraph("Total Jaga dan Jaga Minggu masing-masing berselisih maksimal satu antaranggota."); note.runs[0].font.name="Times New Roman"; note.runs[0].font.size=Pt(10)
    recap=doc.add_table(rows=1,cols=len(summary.columns))
    for column,field in enumerate(summary.columns): shade(recap.cell(0,column),"E7B008"); write_cell(recap.cell(0,column),field,8,True)
    for row in summary.sort_values("Nama").itertuples(index=False):
        cells=recap.add_row().cells
        for column,value in enumerate(row): write_cell(cells[column],str(value),8)
    output=io.BytesIO(); doc.save(output); return output.getvalue()


init_state()
st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@500;600;700;800&display=swap');
:root { --ink:#152a38; --muted:#60717d; --line:#dce5e7; --teal:#0b766e; --teal-dark:#075b55; --mint:#e9f5f2; --paper:#f7faf9; }
.stApp { background:var(--paper); color:var(--ink); font-family:'DM Sans','Helvetica Neue',Arial,sans-serif; }
.block-container { max-width:1180px; padding-top:2.5rem; padding-bottom:4rem; }
h1,h2,h3,[data-testid='stMetricLabel'] { font-family:'Manrope','Helvetica Neue',Arial,sans-serif; color:var(--ink); letter-spacing:-.035em; }
.masthead { border-bottom:1px solid var(--line); padding:0 0 1.85rem; margin-bottom:1.55rem; }
.masthead h1 { margin:.25rem 0 .3rem; font-size:2rem; font-weight:800; }
.masthead p { color:var(--muted); max-width:700px; margin:0; font-size:.96rem; }
.service-line { color:var(--teal); font-family:'Manrope',sans-serif; font-size:.72rem; letter-spacing:.12em; font-weight:800; }
.panel { background:#fff; border:1px solid var(--line); border-radius:10px; padding:1.1rem 1.2rem; margin:.8rem 0 1rem; }
.panel b { font-family:'Manrope',sans-serif; font-size:.95rem; }
[data-testid='stMetric'] { background:#fff; border:1px solid var(--line); border-radius:10px; padding:1rem; box-shadow:none; }
[data-testid='stMetricValue'] { color:var(--teal-dark); font-family:'Manrope',sans-serif; }
div.stButton > button { border-radius:7px; font-family:'DM Sans',sans-serif; font-weight:700; min-height:2.55rem; box-shadow:none; }
div.stButton > button[kind='primary'] { background:var(--teal); border-color:var(--teal); }
div.stButton > button[kind='primary']:hover { background:var(--teal-dark); border-color:var(--teal-dark); }
div[data-testid='stTabs'] button { font-family:'Manrope',sans-serif; font-size:.88rem; }
div[data-testid='stExpander'] { background:#fff; border:1px solid var(--line); border-radius:10px; }
div[data-testid='stDataFrame'] { border:1px solid var(--line); border-radius:10px; overflow:hidden; }
</style>""", unsafe_allow_html=True)
module = st.radio("Modul", ["Penjadwalan Jaga, Review, ERM", "Pembagian Jaga"], horizontal=True, label_visibility="collapsed", key="module")
if module == "Pembagian Jaga":
    render_roster_intake()
    st.stop()
st.markdown("<div class='masthead'><div class='service-line'>DEPARTEMEN BEDAH MULUT & MAKSILOFASIAL • ANGKATAN 15</div><h1>Clinical Duty Roster</h1><p>Susun Jaga, Review, dan ERM dengan distribusi yang tervalidasi. Fairness total dan hari Minggu dikunci sebelum jadwal dapat diekspor.</p></div>", unsafe_allow_html=True)

today=date.today(); mode=st.radio("Jenis jadwal",["Bulanan","Rentang tanggal"],horizontal=True,label_visibility="collapsed",key="schedule_mode")
if mode=="Bulanan":
    left,right=st.columns([2,1])
    with left: month=st.selectbox("Bulan",list(range(1,13)),index=today.month-1,format_func=lambda item:MONTHS[item-1])
    with right: year=int(st.number_input("Tahun",2024,2035,today.year))
    start=date(year,month,1); next_month=date(year+(month==12),1 if month==12 else month+1,1); end=next_month-timedelta(days=1)
else:
    left,right=st.columns(2)
    with left: start=st.date_input("Mulai",value=today.replace(day=1))
    with right: end=st.date_input("Selesai",value=today)

st.markdown("<div class='panel'><b>Komposisi layanan harian</b><br><span style='color:#60717d'>Tetapkan kebutuhan staf untuk periode ini. Tidak ada Kares atau peran Bantu.</span></div>",unsafe_allow_html=True)
c1,c2,c3,c4=st.columns([1,1,1,1.7])
with c1:q_jaga=int(st.number_input("Jaga",0,len(NAMES),4))
with c2:q_review=int(st.number_input("Review",0,len(NAMES),3))
with c3:q_erm=int(st.number_input("ERM",0,len(NAMES),2))
with c4:seed=int(st.number_input("Variasi jadwal",1,999999,1501,help="Ganti angka untuk alternatif yang tetap fair."))

with st.expander("Aturan tim & request khusus"):
    doru=st.multiselect("Doru — tepat 2 orang (tetap Jaga, tidak masuk Review/ERM)",NAMES,default=["Ferrel","Alam"],max_selections=2,key="doru")
    left,right=st.columns(2)
    with left:
        st.markdown("#### Tidak tersedia untuk Jaga")
        request_name=st.selectbox("Nama",NAMES,key="request_name"); request_date=st.date_input("Tanggal",value=start,min_value=start,max_value=end,key="request_date")
        if st.button("Tambahkan request",use_container_width=True): st.session_state.unavailable[request_name].add(request_date.isoformat()); st.rerun()
        requests=[(name,value) for name in NAMES for value in sorted(st.session_state.unavailable[name])]
        if requests:
            st.dataframe(pd.DataFrame(requests,columns=["Nama","Tidak tersedia"]),hide_index=True,use_container_width=True)
            selected=st.selectbox("Hapus request",[f"{name} — {value}" for name,value in requests])
            if st.button("Hapus request terpilih"):
                name,value=selected.split(" — ");st.session_state.unavailable[name].discard(value);st.rerun()
        else: st.caption("Belum ada request khusus.")
    with right:
        st.markdown("#### Larangan pasangan Jaga")
        pair_a=st.selectbox("Nama pertama",NAMES,key="pair_a");pair_b=st.selectbox("Nama kedua",NAMES,index=1,key="pair_b")
        if st.button("Tambah larangan pasangan",use_container_width=True):
            if pair_a==pair_b: st.error("Pilih dua nama berbeda.")
            else: st.session_state.forbidden.add(tuple(sorted((pair_a,pair_b))));st.rerun()
        if st.session_state.forbidden: st.dataframe(pd.DataFrame(sorted(st.session_state.forbidden),columns=["Nama pertama","Nama kedua"]),hide_index=True,use_container_width=True)

total_roles=q_jaga+q_review+q_erm; valid=start<=end and len(doru)==2 and total_roles<=len(NAMES)
if start>end: st.error("Tanggal selesai harus sesudah tanggal mulai.")
elif len(doru)!=2: st.error("Pilih tepat dua Doru.")
elif total_roles>len(NAMES): st.error("Total kuota harian melebihi 11 anggota; satu orang tidak boleh memegang dua peran sehari.")
g1,g2,g3=st.columns([1.3,1,3])
with g1: generate=st.button("Buat jadwal fair",type="primary",use_container_width=True,disabled=not valid,key="generate")
with g2: shuffle=st.button("Acak ulang",use_container_width=True,disabled=not valid,key="shuffle")
with g3: st.caption("Aturan keras: satu peran/hari • Doru hanya Jaga • total Jaga dan Jaga Minggu selisih maksimal satu.")
if generate or shuffle:
    days=[start+timedelta(days=index) for index in range((end-start).days+1)]; quotas={"Jaga":q_jaga,"Review":q_review,"ERM":q_erm}
    with st.spinner("Mencari susunan yang fair untuk semua orang..."):
        schedule,summary,notices=make_schedule(days,quotas,set(doru),st.session_state.unavailable,st.session_state.forbidden,seed+int(shuffle))
    if schedule is None: st.error(" ".join(notices))
    else:
        st.session_state.schedule=schedule;st.session_state.summary=summary;st.session_state.notices=notices;st.session_state.period=(start,end)

if st.session_state.schedule is not None:
    schedule,summary=st.session_state.schedule,st.session_state.summary
    st.markdown("## Hasil jadwal")
    m1,m2,m3,m4=st.columns(4);m1.metric("Hari terjadwal",len(schedule));m2.metric("Rentang total Jaga",f"{summary['Jaga'].min()}–{summary['Jaga'].max()}");m3.metric("Rentang Jaga Minggu",f"{summary['Jaga Minggu'].min()}–{summary['Jaga Minggu'].max()}");m4.metric("Berturut terpanjang",int(summary['Jaga Berturut Terpanjang'].max()))
    tab1,tab2,tab3=st.tabs(["Jadwal","Fairness","Word"])
    with tab1:st.dataframe(schedule[["Tanggal","Jaga","Review","ERM"]],hide_index=True,use_container_width=True,height=520)
    with tab2:
        st.dataframe(summary.sort_values("Nama"),hide_index=True,use_container_width=True)
        st.success("Lolos validasi: total Jaga dan Jaga Minggu masing-masing selisih maksimal satu antaranggota.")
        for notice in st.session_state.notices:st.warning(notice)
    with tab3:
        word_file=schedule_docx(schedule,summary,*st.session_state.period)
        st.download_button("Unduh jadwal Word",word_file,"jadwal-angkatan-15.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",type="primary")
